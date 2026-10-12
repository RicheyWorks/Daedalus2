"""Writers: JSON, Markdown, static HTML dashboard, and xlsx sheets (built from the JSON)."""
import json, html

B = ["in_unreal", "ready", "in_progress", "blender_only", "rejected", "total"]


def write_json(c, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(c, f, indent=1, ensure_ascii=False)


def _mdtable(rows, cols):
    out = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for r in rows:
        out.append("| " + " | ".join(str(r.get(c, "")).replace("|", "/") for c in cols) + " |")
    return "\n".join(out)


def write_md(c, path):
    s = [f"# World census\nGenerated {c['generated']}. Sources: {c['source_counts']}.\n",
         "## Grand total\n" + _mdtable([c["grand_total"]], B),
         "## Per world\n" + _mdtable(c["totals_by_world"], ["world"] + B),
         "## Per owner bot\n" + _mdtable(c["totals_by_owner"], ["owner"] + B),
         "## Per category\n" + _mdtable(c["totals_by_category"], ["category"] + B),
         "## Islands (grid + contest tracker)\n" + _mdtable(c["islands"], ["bot", "island", "ark_origin", "items_reported", "assets_counted", "herb_mechanic", "livestock_food", "water", "in_unreal_reported"]),
         "## Realms\n" + _mdtable(c["realms"], ["Realm", "Status", "Feels like"]),
         "## Island docs\n" + _mdtable(c["island_docs"], ["doc", "title", "water", "livestock", "herbs", "mechanic"]),
         "## World inventory (world x category)\n" + _mdtable(c["inventory"], ["world", "category"] + B)]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(s) + "\n")


def _htable(rows, cols):
    h = "<table><tr>" + "".join(f"<th>{html.escape(x)}</th>" for x in cols) + "</tr>"
    for r in rows:
        h += "<tr>" + "".join(f"<td>{html.escape(str(r.get(x, '')))}</td>" for x in cols) + "</tr>"
    return h + "</table>"


def write_html(c, path):
    top = max([w["total"] for w in c["totals_by_world"]] or [1])
    bars = "".join(f"<div class=b><span>{html.escape(w['world'])}</span><i style='width:{60*w['total']/top:.1f}%'></i><b>{w['total']} ({w['in_unreal']} in UE)</b></div>"
                   for w in c["totals_by_world"])
    g = c["grand_total"]
    doc = f"""<!doctype html><meta charset=utf-8><title>Daedalus world census</title>
<style>body{{font:14px system-ui;margin:24px;background:#0f1a1e;color:#e6eef0}}table{{border-collapse:collapse;margin:8px 0 24px}}
td,th{{border:1px solid #2b4650;padding:3px 8px}}th{{background:#264E58}}.b{{display:flex;gap:8px;align-items:center}}
.b span{{width:260px}}.b i{{display:block;height:12px;background:#4fb3c8}}.k{{display:inline-block;margin-right:24px;font-size:22px}}</style>
<h1>Daedalus world census</h1><p>Generated {html.escape(c['generated'])}</p>
<div>{''.join(f"<span class=k>{g[k]}<br><small>{k}</small></span>" for k in B)}</div>
<h2>Assets per world</h2>{bars}
<h2>Per world</h2>{_htable(c['totals_by_world'], ['world']+B)}
<h2>Islands</h2>{_htable(c['islands'], ['bot','island','ark_origin','items_reported','assets_counted','water','herb_mechanic','livestock_food'])}
<h2>Per owner</h2>{_htable(c['totals_by_owner'], ['owner']+B)}
<h2>Realms</h2>{_htable(c['realms'], ['Realm','Status','Feels like'])}
<h2>World x category</h2>{_htable(c['inventory'], ['world','category']+B)}"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(doc)


def add_xlsx_sheets(c, xlsx_path, prefix=""):
    """Add/replace census sheets in an existing workbook (creates it if missing). Other sheets are untouched."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    try:
        wb = load_workbook(xlsx_path)
    except FileNotFoundError:
        wb = Workbook(); wb.remove(wb.active)
    sheets = [("World Inventory", c["inventory"], ["world", "category"] + B),
              ("Asset Catalog", c["assets"], ["name", "kind", "category", "world", "owner", "status", "queue_status", "blender_files", "uassets", "path"]),
              ("Totals", [dict(scope="GRAND TOTAL", key="all", **c["grand_total"])] +
               [dict(scope="world", key=r["world"], **r) for r in c["totals_by_world"]] +
               [dict(scope="owner", key=r["owner"], **r) for r in c["totals_by_owner"]] +
               [dict(scope="category", key=r["category"], **r) for r in c["totals_by_category"]], ["scope", "key"] + B),
              ("Islands", c["islands"], list(c["islands"][0].keys()) if c["islands"] else ["bot"]),
              ("Realms", c["realms"], ["Realm", "Status", "Feels like", "What's in it", "ADR", "doc"])]
    for title, rows, cols in sheets:
        t = (prefix + title)[:31]
        if t in wb.sheetnames:
            del wb[t]
        ws = wb.create_sheet(t); ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="264E58")
        for r in rows:
            ws.append([r.get(k, "") for k in cols])
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for i, k in enumerate(cols):
            ws.column_dimensions[chr(65 + i)].width = 34 if k in ("world", "name", "key", "island", "path", "Realm") else 14
    wb.save(xlsx_path)
